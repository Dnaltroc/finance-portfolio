from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from finlib.models.variance import (
    Thresholds, commentary, compute_variance, pnl_summary, rollup, top_flags, validate_ledger,
)

SD = Path(__file__).resolve().parents[1] / "projects" / "variance" / "sample_data"


def _mini():
    actuals = pd.DataFrame([
        dict(period="2026-09", department="Sales", account="Product revenue", account_type="revenue", amount=110),
        dict(period="2026-09", department="Sales", account="Salaries", account_type="expense", amount=60),
        dict(period="2026-09", department="Ops", account="Rent", account_type="expense", amount=100),
    ])
    budget = pd.DataFrame([
        dict(period="2026-09", department="Sales", account="Product revenue", account_type="revenue", amount=100),
        dict(period="2026-09", department="Sales", account="Salaries", account_type="expense", amount=50),
        dict(period="2026-09", department="Ops", account="Rent", account_type="expense", amount=100),
    ])
    return validate_ledger(actuals, "a"), validate_ledger(budget, "b")


def test_sign_logic_revenue_vs_expense():
    a, b = _mini()
    d = compute_variance(a, b, Thresholds(abs_amount=1, pct=0.05))
    rev = d[d["account"] == "Product revenue"].iloc[0]
    exp = d[d["account"] == "Salaries"].iloc[0]
    flat = d[d["account"] == "Rent"].iloc[0]
    assert rev["variance_amt"] == 10 and rev["fav_unfav"] == "F"
    assert exp["variance_amt"] == 10 and exp["fav_unfav"] == "U"   # same raw sign, opposite judgment
    assert flat["fav_unfav"] == "-"


def test_threshold_either_rule_flags():
    a, b = _mini()
    d = compute_variance(a, b, Thresholds(abs_amount=1_000_000, pct=0.15))  # only pct can trip
    assert d.set_index("account").loc["Salaries", "flag"]            # 20%
    assert not d.set_index("account").loc["Product revenue", "flag"] # 10%
    d2 = compute_variance(a, b, Thresholds(abs_amount=10, pct=0.99))  # only abs can trip
    assert d2.set_index("account").loc["Product revenue", "flag"]


def test_rollups_tie_to_detail():
    a, b = _mini()
    d = compute_variance(a, b, Thresholds())
    by_dept = rollup(d, ["department"]).set_index("department")
    # Rollup is on contribution: Sales has revenue 110 and expense 60 -> 50 actual, 50 budget
    assert by_dept.loc["Sales", "actual"] == 110 - 60 and by_dept.loc["Sales", "budget"] == 100 - 50
    assert by_dept.loc["Sales", "variance_amt"] == 0 and by_dept.loc["Sales", "fav_unfav"] == "-"
    assert by_dept["variance_amt"].sum() == d["favorable_amt"].sum()
    s = pnl_summary(d)
    net = s.set_index("line")
    assert net.loc["Net", "actual"] == 110 - 160
    assert net.loc["Net", "variance_amt"] == (110 - 100) - (160 - 150)


def test_blank_amount_is_rejected_not_zeroed():
    a, b = _mini()
    a.loc[0, "amount"] = float("nan")
    with pytest.raises(ValueError, match="row 2"):
        validate_ledger(a, "actuals")


def test_validation_rejects_bad_rows_with_row_numbers():
    bad = pd.DataFrame([dict(period="2026-9", department="X", account="Y", account_type="expense", amount=1)])
    with pytest.raises(ValueError, match="row 2"):
        validate_ledger(bad, "actuals")
    bad2 = pd.DataFrame([dict(period="2026-09", department="X", account="Y", account_type="widget", amount=1)])
    with pytest.raises(ValueError, match="account_type"):
        validate_ledger(bad2, "actuals")


def test_missing_budget_line_shows_as_full_variance():
    a, b = _mini()
    a = pd.concat([a, pd.DataFrame([dict(period="2026-09", department="Ops", account="Unbudgeted",
                                         account_type="expense", amount=40)])])
    d = compute_variance(a, b, Thresholds(abs_amount=1))
    row = d[d["account"] == "Unbudgeted"].iloc[0]
    assert row["budget"] == 0 and row["variance_amt"] == 40 and row["fav_unfav"] == "U"
    assert pd.isna(row["variance_pct"])  # no % on a zero budget


def test_commentary_one_bullet_per_flag():
    a, b = _mini()
    d = compute_variance(a, b, Thresholds(abs_amount=5, pct=0.5))
    flags = top_flags(d, 10)
    text = commentary(flags, "2026-09")
    assert text.count("\n- ") == len(flags) == 2
    assert "driven by ___" in text


def test_sample_data_runs_end_to_end(tmp_path: Path):
    from projects.variance import run_variance

    cfg = SD.parent / "config.yaml"
    for period in ("2026-08", "2026-09"):
        xlsx, md = run_variance(SD / "actuals.csv", SD / "budget.csv", period, cfg, tmp_path)
        wb = load_workbook(xlsx)
        assert wb.sheetnames == ["Summary", "Flags", "Detail", "Inputs"]
        assert md.read_text().startswith("# Variance commentary")
    ws = load_workbook(tmp_path / "variance_2026-09.xlsx")["Flags"]
    accounts = [ws.cell(row=r, column=2).value for r in range(4, 8)]
    assert "Paid media" in accounts  # the planted overspend is in the top flags
