from pathlib import Path

from openpyxl import load_workbook

from projects.comps import run_comps


def test_workbook_has_three_sheets_and_summary_rows(tmp_path: Path):
    path = run_comps(["AAPL", "MSFT", "GOOGL"], "test", tmp_path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Comps", "Inputs", "Notes"]
    ws = wb["Comps"]
    headers = [c.value for c in ws[4]]
    assert headers[0] == "Ticker" and "EV / EBITDA" in headers
    labels = [ws.cell(row=r, column=2).value for r in range(5, 10)]
    assert labels[:3] == ["Apple Inc.", "Microsoft Corporation", "Alphabet Inc."]
    assert labels[3:] == ["Median", "Mean"]
    # Multiples are stored as numbers with an "x" number format, not as strings.
    col = headers.index("EV / EBITDA") + 1
    cell = ws.cell(row=5, column=col)
    assert isinstance(cell.value, float) and cell.number_format.endswith('"NM"')


def test_inputs_sheet_cites_a_source_per_value(tmp_path: Path):
    path = run_comps(["AAPL"], "test", tmp_path)
    ws = load_workbook(path)["Inputs"]
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    assert len(rows) == 9  # price, shares, 5 flow concepts, 2 stock concepts
    assert all(r[4] for r in rows)  # source column populated
