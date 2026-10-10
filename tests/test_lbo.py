import shutil
import subprocess
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from finlib.excel_writer.lbo_workbook import write_lbo_workbook
from finlib.models.lbo import LBOAssumptions, irr, run_lbo, sensitivity
from projects.lbo import load_deal

DEAL = Path(__file__).resolve().parents[1] / "projects" / "lbo" / "deal.yaml"
SOFFICE = shutil.which("soffice")


def _recalc(path: Path, out_dir: Path) -> Path:
    """Recalculate every formula with LibreOffice headless; returns the recalculated copy."""
    subprocess.run([SOFFICE, "--headless", "--convert-to", "xlsx", "--outdir", str(out_dir), str(path)],
                   check=True, capture_output=True, timeout=180)
    return out_dir / path.name


def test_irr_matches_closed_form():
    # -100 now, +121 in two years is exactly 10%
    assert abs(irr([-100, 0, 121]) - 0.10) < 1e-9
    with pytest.raises(ValueError):
        irr([-1, -2, -3])


def test_sources_equal_uses():
    a = load_deal(DEAL)
    uses = a.entry_ev + a.fees + a.min_cash
    sources = a.total_debt + a.sponsor_equity
    assert abs(uses - sources) < 1e-9


def test_debt_schedule_identities():
    a = load_deal(DEAL)
    r = run_lbo(a)
    for name in r.tranche_end:
        for y in range(1, a.hold_years + 1):
            assert r.tranche_begin[name][y] == r.tranche_end[name][y - 1]
            assert abs(r.tranche_end[name][y] - (r.tranche_begin[name][y] - r.tranche_mandatory[name][y]
                                                - r.tranche_sweep[name][y])) < 1e-9
            assert r.tranche_end[name][y] >= -1e-9  # never over-repay
    # Senior Notes has sweep=False, so it only amortizes per its amort_pct (0 here)
    assert np.allclose(r.tranche_end["Senior Notes"], r.tranche_end["Senior Notes"][0])
    # Cash never drops below the minimum once the sweep runs
    assert (r.cash_balance[1:] >= a.min_cash - 1e-9).all()


def test_moic_and_irr_consistent():
    a = load_deal(DEAL)
    r = run_lbo(a)
    assert abs(r.moic - r.exit_equity / r.sponsor_equity) < 1e-12
    assert abs(r.irr - irr(r.equity_cashflows)) < 1e-12
    assert 0.15 < r.irr < 0.30  # sanity band for a 9x in / 9x out deal at 5.5x leverage


def test_sensitivity_monotonic():
    a = load_deal(DEAL)
    g = sensitivity(a, [8.0, 9.0, 10.0], [8.0, 9.0, 10.0])
    assert (np.diff(g, axis=1) > 0).all()   # higher exit multiple -> higher IRR
    assert (np.diff(g, axis=0) < 0).all()   # higher entry multiple -> lower IRR


def test_overlevered_deal_flags_shortfall_and_wiped_equity():
    a = load_deal(DEAL)
    b = a.model_copy(deep=True)
    b.tranches[0].ebitda_multiple = 7.0
    b.tranches[0].rate = 0.14
    b.tranches[0].amort_pct = 0.10
    r = run_lbo(b, "Downside")
    assert r.cash_shortfall_years, "cash should fall below minimum somewhere"
    assert r.exit_equity == 0.0 and r.moic == 0.0 and not r.irr_defined
    assert not np.isnan(run_lbo(a).irr)


def test_sweep_stops_at_zero_balance():
    a = load_deal(DEAL)
    b = a.model_copy(deep=True)
    b.tranches[0].ebitda_multiple = 0.3  # tiny TLB, fully repaid by the sweep early
    r = run_lbo(b)
    tlb = r.tranche_end["Term Loan B"]
    assert (tlb >= -1e-9).all() and tlb[-1] == 0.0
    assert (r.cash_balance[1:] >= b.min_cash - 1e-9).all()  # excess sits in cash once debt is gone


def test_scenario_validation():
    raw = {"entry_revenue": 100, "entry_ebitda": 20, "entry_multiple": 8, "tranches": [],
           "scenarios": {"Base": {"revenue_growth": [0.05, 0.05], "ebitda_margin": [0.2], "exit_multiple": 8}}}
    with pytest.raises(ValueError):
        LBOAssumptions(**raw)


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed; cannot recalc formulas")
def test_excel_irr_matches_python_within_tolerance(tmp_path: Path):
    a = load_deal(DEAL)
    path = write_lbo_workbook(a, tmp_path)
    calc = _recalc(path, tmp_path / "calc")
    ws = load_workbook(calc, data_only=True)["Returns"]
    py = run_lbo(a)
    assert abs(ws["C16"].value - py.irr) < 0.001
    assert abs(ws["C17"].value - py.moic) < 0.001


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed; cannot recalc formulas")
def test_scenario_dropdown_switches_every_sheet(tmp_path: Path):
    a = load_deal(DEAL)
    path = write_lbo_workbook(a, tmp_path)
    wb = load_workbook(path)
    wb["Assumptions"]["C4"] = "Downside"
    switched = tmp_path / "downside.xlsx"
    wb.save(switched)
    calc = _recalc(switched, tmp_path / "calc")
    ws = load_workbook(calc, data_only=True)["Returns"]
    py = run_lbo(a, "Downside")
    assert abs(ws["C16"].value - py.irr) < 0.001
    assert abs(ws["C17"].value - py.moic) < 0.001
