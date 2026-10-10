"""House style for every workbook: banker conventions, nothing fancy."""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")  # navy
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
BODY_FONT = Font(name="Calibri", size=10)
INPUT_FONT = Font(name="Calibri", size=10, color="0000FF")  # blue = hardcoded input
SUMMARY_FONT = Font(name="Calibri", size=10, bold=True)
SUMMARY_FILL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
TOP_BORDER = Border(top=Side(style="thin", color="000000"))
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right")

NUM_FMT = '#,##0;(#,##0);"-"'          # $ millions, parentheses for negatives
X_FMT = '0.0"x";(0.0"x");"NM"'          # multiples
PCT_FMT = '0.0%;(0.0%);"-"'
PRICE_FMT = '#,##0.00'


def autosize(ws, min_width: int = 8, max_width: int = 42) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[letter].width = max(min_width, min(max_width, width + 2))
