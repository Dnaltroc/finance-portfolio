"""LBO workbook with live formulas.

Every computed cell is an Excel formula pointing at the Assumptions sheet, so a reviewer
can change an input or the scenario dropdown and watch IRR move. The Python engine in
finlib/models/lbo.py is the reference; tests recalc this workbook with LibreOffice and
assert the Excel IRR matches Python within 0.1%.

Sheet map
  Assumptions  inputs (blue), scenario selector (data validation), active scenario via INDEX/MATCH
  Operating    revenue -> FCF, years 0..N across columns
  Debt         per-tranche schedule with mandatory amortization and cash sweep, cash roll
  Returns      exit bridge, MOIC, IRR(), equity cash flows
  Sensitivity  IRR grid from the Python engine (values; Excel data tables are not writable via openpyxl)
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from finlib.models.lbo import LBOAssumptions, sensitivity
from .styles import BODY_FONT, HEADER_FILL, HEADER_FONT, INPUT_FONT, SUMMARY_FILL, SUMMARY_FONT

MONEY = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
X = '0.00"x"'
BOLD = Font(name="Calibri", size=10, bold=True)
FIRST_YEAR_COL = 3  # column C = year 0


def _yc(year: int) -> str:
    return L(FIRST_YEAR_COL + year)


def _label(ws, row: int, text: str, bold: bool = False) -> None:
    ws.cell(row=row, column=1, value=text).font = BOLD if bold else BODY_FONT


def _inp(ws, row: int, col: int, value, fmt: str | None = None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = INPUT_FONT
    if fmt:
        c.number_format = fmt
    return c


def _f(ws, row: int, col: int, formula: str, fmt: str | None = None, bold: bool = False):
    c = ws.cell(row=row, column=col, value=formula)
    c.font = BOLD if bold else BODY_FONT
    if fmt:
        c.number_format = fmt
    return c


def _year_header(ws, row: int, n: int, label: str = "Year") -> None:
    ws.cell(row=row, column=1, value=label).font = HEADER_FONT
    ws.cell(row=row, column=1).fill = HEADER_FILL
    ws.cell(row=row, column=2).fill = HEADER_FILL
    for y in range(n + 1):
        c = ws.cell(row=row, column=FIRST_YEAR_COL + y, value=y)
        c.font, c.fill = HEADER_FONT, HEADER_FILL


def write_lbo_workbook(a: LBOAssumptions, out_dir: Path,
                       entry_grid=(7.0, 8.0, 9.0, 10.0, 11.0),
                       exit_grid=(7.0, 8.0, 9.0, 10.0, 11.0)) -> Path:
    n = a.hold_years
    wb = Workbook()
    A = _assumptions(wb.active, a)
    O = _operating(wb.create_sheet("Operating"), a, A)
    D = _debt(wb.create_sheet("Debt"), a, A, O)
    # Operating interest and taxes depend on Debt; fill those rows now that D exists.
    _link_interest(wb["Operating"], a, O, D)
    _returns(wb.create_sheet("Returns"), a, A, O, D)
    _sensitivity(wb.create_sheet("Sensitivity"), a, list(entry_grid), list(exit_grid))
    for ws in wb.worksheets:
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 12
        for y in range(n + 2):
            ws.column_dimensions[L(FIRST_YEAR_COL + y)].width = 12
    out_dir.mkdir(parents=True, exist_ok=True)
    safe = "".join(ch if ch.isalnum() else "_" for ch in a.deal_name)
    path = out_dir / f"lbo_{safe}.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------- Assumptions ---------
def _assumptions(ws, a: LBOAssumptions) -> dict:
    ws.title = "Assumptions"
    n = a.hold_years
    ws["A1"] = f"LBO Assumptions: {a.deal_name}"
    ws["A1"].font = SUMMARY_FONT
    ws["A2"] = "$ in millions. Blue = input. Change the scenario in C4 and every sheet recalculates."
    ws["A2"].font = BODY_FONT

    ref: dict = {}
    _label(ws, 4, "Active scenario", bold=True)
    names = list(a.scenarios.keys())
    sel = _inp(ws, 4, 3, a.active_scenario)
    dv = DataValidation(type="list", formula1='"' + ",".join(names) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(sel)
    ref["scenario"] = "Assumptions!$C$4"

    rows = [
        ("entry_revenue", "Entry revenue", a.entry_revenue, MONEY),
        ("entry_ebitda", "Entry EBITDA", a.entry_ebitda, MONEY),
        ("entry_multiple", "Entry EV / EBITDA", a.entry_multiple, X),
        ("fees_pct_ev", "Transaction fees (% of EV)", a.fees_pct_ev, PCT),
        ("min_cash", "Minimum cash", a.min_cash, MONEY),
        ("tax_rate", "Tax rate", a.tax_rate, PCT),
        ("da_pct_revenue", "D&A (% of revenue)", a.da_pct_revenue, PCT),
        ("capex_pct_revenue", "Capex (% of revenue)", a.capex_pct_revenue, PCT),
        ("nwc_pct_revenue", "NWC (% of revenue)", a.nwc_pct_revenue, PCT),
    ]
    r = 6
    _label(ws, r, "Deal inputs", bold=True); r += 1
    for key, label, val, fmt in rows:
        _label(ws, r, label)
        _inp(ws, r, 3, val, fmt)
        ref[key] = f"Assumptions!$C${r}"
        r += 1

    # Scenario table: col C = active (formula), D.. = each scenario (inputs)
    r += 1
    _label(ws, r, "Scenario inputs", bold=True)
    hdr = r + 1
    ws.cell(row=hdr, column=1, value="Driver").font = HEADER_FONT
    ws.cell(row=hdr, column=1).fill = HEADER_FILL
    ws.cell(row=hdr, column=2, value="Year").font = HEADER_FONT
    ws.cell(row=hdr, column=2).fill = HEADER_FILL
    ws.cell(row=hdr, column=3, value="Active").font = HEADER_FONT
    ws.cell(row=hdr, column=3).fill = HEADER_FILL
    for j, name in enumerate(names):
        c = ws.cell(row=hdr, column=4 + j, value=name)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
    first_sc_col, last_sc_col = L(4), L(3 + len(names))
    hdr_range = f"$ {first_sc_col}${hdr}:${last_sc_col}${hdr}".replace(" ", "")

    def active_formula(row: int) -> str:
        return (f"=INDEX(${first_sc_col}{row}:${last_sc_col}{row},1,"
                f"MATCH({ref['scenario']},{hdr_range},0))")

    r = hdr + 1
    ref["growth"], ref["margin"] = {}, {}
    for y in range(1, n + 1):
        _label(ws, r, "Revenue growth")
        ws.cell(row=r, column=2, value=y).font = BODY_FONT
        _f(ws, r, 3, active_formula(r), PCT)
        for j, name in enumerate(names):
            _inp(ws, r, 4 + j, a.scenarios[name].revenue_growth[y - 1], PCT)
        ref["growth"][y] = f"Assumptions!$C${r}"
        r += 1
    for y in range(1, n + 1):
        _label(ws, r, "EBITDA margin")
        ws.cell(row=r, column=2, value=y).font = BODY_FONT
        _f(ws, r, 3, active_formula(r), PCT)
        for j, name in enumerate(names):
            _inp(ws, r, 4 + j, a.scenarios[name].ebitda_margin[y - 1], PCT)
        ref["margin"][y] = f"Assumptions!$C${r}"
        r += 1
    _label(ws, r, "Exit EV / EBITDA")
    _f(ws, r, 3, active_formula(r), X)
    for j, name in enumerate(names):
        _inp(ws, r, 4 + j, a.scenarios[name].exit_multiple, X)
    ref["exit_multiple"] = f"Assumptions!$C${r}"
    r += 2

    # Tranches
    _label(ws, r, "Debt tranches", bold=True); r += 1
    heads = ["Tranche", "x EBITDA", "Rate", "Amort %/yr", "Sweep (1/0)", "Amount"]
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=r, column=j, value=h)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
    r += 1
    ref["tranches"] = {}
    for t in a.tranches:
        ws.cell(row=r, column=1, value=t.name).font = BODY_FONT
        _inp(ws, r, 2, t.ebitda_multiple, X)
        _inp(ws, r, 3, t.rate, PCT)
        _inp(ws, r, 4, t.amort_pct, PCT)
        _inp(ws, r, 5, 1 if t.sweep else 0)
        _f(ws, r, 6, f"=B{r}*{ref['entry_ebitda']}", MONEY)
        ref["tranches"][t.name] = dict(mult=f"Assumptions!$B${r}", rate=f"Assumptions!$C${r}",
                                       amort=f"Assumptions!$D${r}", sweep=f"Assumptions!$E${r}",
                                       amount=f"Assumptions!$F${r}")
        r += 1
    first_t, last_t = r - len(a.tranches), r - 1
    r += 1

    # Sources and uses
    _label(ws, r, "Sources and uses", bold=True); r += 1
    _label(ws, r, "Entry enterprise value"); _f(ws, r, 3, f"={ref['entry_ebitda']}*{ref['entry_multiple']}", MONEY)
    ref["entry_ev"] = f"Assumptions!$C${r}"; r += 1
    _label(ws, r, "Transaction fees"); _f(ws, r, 3, f"={ref['entry_ev']}*{ref['fees_pct_ev']}", MONEY)
    ref["fees"] = f"Assumptions!$C${r}"; r += 1
    _label(ws, r, "Cash to balance sheet"); _f(ws, r, 3, f"={ref['min_cash']}", MONEY); r += 1
    _label(ws, r, "Total uses", bold=True); _f(ws, r, 3, f"=SUM(C{r-3}:C{r-1})", MONEY, bold=True)
    ref["total_uses"] = f"Assumptions!$C${r}"; r += 1
    _label(ws, r, "Total debt"); _f(ws, r, 3, f"=SUM(F{first_t}:F{last_t})", MONEY)
    ref["total_debt"] = f"Assumptions!$C${r}"; r += 1
    _label(ws, r, "Sponsor equity", bold=True); _f(ws, r, 3, f"={ref['total_uses']}-{ref['total_debt']}", MONEY, bold=True)
    ref["sponsor_equity"] = f"Assumptions!$C${r}"; r += 1
    _label(ws, r, "Leverage (debt / EBITDA)"); _f(ws, r, 3, f"={ref['total_debt']}/{ref['entry_ebitda']}", X)
    return ref


# ---------------------------------------------------------------- Operating -----------
def _operating(ws, a: LBOAssumptions, A: dict) -> dict:
    n = a.hold_years
    ws["A1"] = "Operating build"; ws["A1"].font = SUMMARY_FONT
    _year_header(ws, 3, n)
    O = {}
    rows = ["revenue", "growth", "margin", "ebitda", "da", "ebit", "interest", "taxes",
            "capex", "dnwc", "fcf_pre", "fcf_post"]
    labels = {"revenue": "Revenue", "growth": "  growth", "margin": "  EBITDA margin", "ebitda": "EBITDA",
              "da": "D&A", "ebit": "EBIT", "interest": "Cash interest (from Debt)",
              "taxes": "Taxes", "capex": "Capex", "dnwc": "Change in NWC",
              "fcf_pre": "FCF before interest and debt service", "fcf_post": "FCF after interest"}
    for i, k in enumerate(rows):
        O[k] = 4 + i
        _label(ws, O[k], labels[k], bold=k in ("revenue", "ebitda", "fcf_post"))

    for y in range(n + 1):
        c = _yc(y)
        p = _yc(y - 1) if y > 0 else None
        if y == 0:
            _f(ws, O["revenue"], FIRST_YEAR_COL, f"={A['entry_revenue']}", MONEY, bold=True)
            _f(ws, O["margin"], FIRST_YEAR_COL, f"={A['entry_ebitda']}/{A['entry_revenue']}", PCT)
        else:
            _f(ws, O["growth"], FIRST_YEAR_COL + y, f"={A['growth'][y]}", PCT)
            _f(ws, O["revenue"], FIRST_YEAR_COL + y, f"={p}{O['revenue']}*(1+{c}{O['growth']})", MONEY, bold=True)
            _f(ws, O["margin"], FIRST_YEAR_COL + y, f"={A['margin'][y]}", PCT)
        col = FIRST_YEAR_COL + y
        _f(ws, O["ebitda"], col, f"={c}{O['revenue']}*{c}{O['margin']}", MONEY, bold=True)
        _f(ws, O["da"], col, f"={c}{O['revenue']}*{A['da_pct_revenue']}", MONEY)
        _f(ws, O["ebit"], col, f"={c}{O['ebitda']}-{c}{O['da']}", MONEY)
        _f(ws, O["capex"], col, f"={c}{O['revenue']}*{A['capex_pct_revenue']}", MONEY)
        if y == 0:
            for k in ("dnwc", "interest", "taxes", "fcf_pre", "fcf_post"):
                _f(ws, O[k], col, "=0", MONEY)
        else:
            _f(ws, O["dnwc"], col, f"=({c}{O['revenue']}-{p}{O['revenue']})*{A['nwc_pct_revenue']}", MONEY)
            _f(ws, O["taxes"], col, f"=MAX(0,{c}{O['ebit']}-{c}{O['interest']})*{A['tax_rate']}", MONEY)
            _f(ws, O["fcf_pre"], col, f"={c}{O['ebitda']}-{c}{O['taxes']}-{c}{O['capex']}-{c}{O['dnwc']}", MONEY)
            _f(ws, O["fcf_post"], col, f"={c}{O['fcf_pre']}-{c}{O['interest']}", MONEY, bold=True)
    return O


def _link_interest(ws, a: LBOAssumptions, O: dict, D: dict) -> None:
    for y in range(1, a.hold_years + 1):
        c = _yc(y)
        _f(ws, O["interest"], FIRST_YEAR_COL + y, f"=Debt!{c}{D['total_interest']}", MONEY)


# ---------------------------------------------------------------- Debt ----------------
def _debt(ws, a: LBOAssumptions, A: dict, O: dict) -> dict:
    n = a.hold_years
    ws["A1"] = "Debt schedule and cash. Interest on beginning balance; sweep senior-first after mandatory amortization."
    ws["A1"].font = SUMMARY_FONT
    _year_header(ws, 3, n)
    D: dict = {"tranche": {}}
    r = 4
    for t in a.tranches:
        T = A["tranches"][t.name]
        _label(ws, r, t.name, bold=True); r += 1
        rows = {"begin": "Beginning balance", "mand": "Mandatory amortization", "sweep": "Cash sweep",
                "end": "Ending balance", "interest": "Interest"}
        D["tranche"][t.name] = {k: r + i for i, k in enumerate(rows)}
        for i, (k, lab) in enumerate(rows.items()):
            _label(ws, r + i, "  " + lab, bold=(k == "end"))
        r += len(rows) + 1
    # Cash roll
    _label(ws, r, "Cash", bold=True); r += 1
    cash_rows = {"cash_begin": "Beginning cash", "fcf_post": "FCF after interest (from Operating)",
                 "mand_total": "Less: mandatory amortization", "pre_sweep": "Cash before sweep",
                 "available": "Available for sweep (above minimum)", "sweep_total": "Less: cash sweep",
                 "cash_end": "Ending cash"}
    for i, (k, lab) in enumerate(cash_rows.items()):
        D[k] = r + i
        _label(ws, r + i, "  " + lab, bold=(k == "cash_end"))
    r += len(cash_rows) + 1
    _label(ws, r, "Total debt", bold=True); D["total_debt"] = r; r += 1
    _label(ws, r, "Total interest", bold=True); D["total_interest"] = r; r += 1
    _label(ws, r, "Net debt", bold=True); D["net_debt"] = r; r += 1
    _label(ws, r, "Net debt / EBITDA"); D["leverage"] = r; r += 1
    _label(ws, r, "Cash below minimum?", bold=True); D["shortfall"] = r

    names = [t.name for t in a.tranches]
    for y in range(n + 1):
        col, c = FIRST_YEAR_COL + y, _yc(y)
        p = _yc(y - 1) if y > 0 else None
        prev_sweeps = []
        for t in a.tranches:
            T, R = A["tranches"][t.name], D["tranche"][t.name]
            if y == 0:
                _f(ws, R["begin"], col, "=0", MONEY)
                _f(ws, R["mand"], col, "=0", MONEY)
                _f(ws, R["sweep"], col, "=0", MONEY)
                _f(ws, R["end"], col, f"={T['amount']}", MONEY, bold=True)
                _f(ws, R["interest"], col, "=0", MONEY)
            else:
                _f(ws, R["begin"], col, f"={p}{R['end']}", MONEY)
                _f(ws, R["mand"], col, f"=MIN({c}{R['begin']},{T['amount']}*{T['amort']})", MONEY)
                taken = "+".join(f"{c}{s}" for s in prev_sweeps) or "0"
                _f(ws, R["sweep"], col,
                   f"=IF({T['sweep']}=1,MAX(0,MIN({c}{D['available']}-({taken}),{c}{R['begin']}-{c}{R['mand']})),0)",
                   MONEY)
                _f(ws, R["end"], col, f"={c}{R['begin']}-{c}{R['mand']}-{c}{R['sweep']}", MONEY, bold=True)
                _f(ws, R["interest"], col, f"={c}{R['begin']}*{T['rate']}", MONEY)
                prev_sweeps.append(R["sweep"])
        mand_sum = "+".join(f"{c}{D['tranche'][k]['mand']}" for k in names)
        sweep_sum = "+".join(f"{c}{D['tranche'][k]['sweep']}" for k in names)
        end_sum = "+".join(f"{c}{D['tranche'][k]['end']}" for k in names)
        int_sum = "+".join(f"{c}{D['tranche'][k]['interest']}" for k in names)
        if y == 0:
            _f(ws, D["cash_begin"], col, "=0", MONEY)
            _f(ws, D["fcf_post"], col, "=0", MONEY)
            _f(ws, D["mand_total"], col, "=0", MONEY)
            _f(ws, D["pre_sweep"], col, f"={A['min_cash']}", MONEY)
            _f(ws, D["available"], col, "=0", MONEY)
            _f(ws, D["sweep_total"], col, "=0", MONEY)
            _f(ws, D["cash_end"], col, f"={A['min_cash']}", MONEY, bold=True)
        else:
            _f(ws, D["cash_begin"], col, f"={p}{D['cash_end']}", MONEY)
            _f(ws, D["fcf_post"], col, f"=Operating!{c}{O['fcf_post']}", MONEY)
            _f(ws, D["mand_total"], col, f"={mand_sum}", MONEY)
            _f(ws, D["pre_sweep"], col, f"={c}{D['cash_begin']}+{c}{D['fcf_post']}-{c}{D['mand_total']}", MONEY)
            _f(ws, D["available"], col, f"=MAX(0,{c}{D['pre_sweep']}-{A['min_cash']})", MONEY)
            _f(ws, D["sweep_total"], col, f"={sweep_sum}", MONEY)
            _f(ws, D["cash_end"], col, f"={c}{D['pre_sweep']}-{c}{D['sweep_total']}", MONEY, bold=True)
        _f(ws, D["total_debt"], col, f"={end_sum}", MONEY, bold=True)
        _f(ws, D["total_interest"], col, f"={int_sum}", MONEY, bold=True)
        _f(ws, D["net_debt"], col, f"={c}{D['total_debt']}-{c}{D['cash_end']}", MONEY, bold=True)
        _f(ws, D["leverage"], col, f"={c}{D['net_debt']}/Operating!{c}{O['ebitda']}", X)
        _f(ws, D["shortfall"], col, f'=IF({c}{D["cash_end"]}<{A["min_cash"]}-0.001,"SHORTFALL","")', bold=True)
    return D


# ---------------------------------------------------------------- Returns -------------
def _returns(ws, a: LBOAssumptions, A: dict, O: dict, D: dict) -> None:
    n = a.hold_years
    xc = _yc(n)
    ws["A1"] = "Returns"; ws["A1"].font = SUMMARY_FONT
    lines = [
        ("Exit year EBITDA", f"=Operating!{xc}{O['ebitda']}", MONEY),
        ("Exit EV / EBITDA", f"={A['exit_multiple']}", X),
        ("Exit enterprise value", "=C3*C4", MONEY),
        ("Less: exit debt", f"=Debt!{xc}{D['total_debt']}", MONEY),
        ("Plus: exit cash", f"=Debt!{xc}{D['cash_end']}", MONEY),
        ("Exit equity value", "=MAX(0,C5-C6+C7)", MONEY),
        ("Sponsor equity invested", f"={A['sponsor_equity']}", MONEY),
        ("MOIC", "=C8/C9", X),
    ]
    for i, (lab, f, fmt) in enumerate(lines, start=3):
        _label(ws, i, lab, bold=lab in ("Exit equity value", "MOIC"))
        _f(ws, i, 3, f, fmt, bold=lab in ("Exit equity value", "MOIC"))
    _year_header(ws, 13, n)
    _label(ws, 14, "Sponsor equity cash flows", bold=True)
    for y in range(n + 1):
        col = FIRST_YEAR_COL + y
        if y == 0:
            _f(ws, 14, col, "=-C9", MONEY)
        elif y == n:
            _f(ws, 14, col, "=C8", MONEY)
        else:
            _f(ws, 14, col, "=0", MONEY)
    _label(ws, 16, "IRR", bold=True)
    c = _f(ws, 16, 3, f"=IRR(C14:{xc}14)", PCT, bold=True)
    c.fill = SUMMARY_FILL
    _label(ws, 17, "MOIC", bold=True)
    _f(ws, 17, 3, "=C10", X, bold=True).fill = SUMMARY_FILL


# ---------------------------------------------------------------- Sensitivity ---------
def _sensitivity(ws, a: LBOAssumptions, entry_grid: list[float], exit_grid: list[float]) -> None:
    ws["A1"] = f"IRR sensitivity: entry multiple (rows) vs exit multiple (columns), {a.active_scenario} scenario"
    ws["A1"].font = SUMMARY_FONT
    ws["A2"] = ("Values computed by the Python engine at export time, not live formulas "
                "(Excel data tables cannot be written by openpyxl). Re-run the CLI after changing inputs.")
    ws["A2"].font = BODY_FONT
    grid = sensitivity(a, entry_grid, exit_grid)
    ws.cell(row=4, column=1, value="Entry \\ Exit").font = HEADER_FONT
    ws.cell(row=4, column=1).fill = HEADER_FILL
    for j, xm in enumerate(exit_grid):
        c = ws.cell(row=4, column=2 + j, value=xm)
        c.font, c.fill, c.number_format = HEADER_FONT, HEADER_FILL, X
    for i, em in enumerate(entry_grid):
        c = ws.cell(row=5 + i, column=1, value=em)
        c.font, c.number_format = BOLD, X
        for j in range(len(exit_grid)):
            v = grid[i, j]
            cell = ws.cell(row=5 + i, column=2 + j, value=None if np.isnan(v) else float(v))
            cell.number_format, cell.font = PCT, BODY_FONT
    rng = f"B5:{L(1 + len(exit_grid))}{4 + len(entry_grid)}"
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="FFC7CE",
                                                      mid_type="percentile", mid_value=50, mid_color="FFFFFF",
                                                      end_type="max", end_color="C6EFCE"))
    ws.column_dimensions["A"].width = 14
